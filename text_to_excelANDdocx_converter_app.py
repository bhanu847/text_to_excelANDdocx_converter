from flask import Flask, request, render_template_string, redirect, flash, send_file, url_for
from werkzeug.utils import secure_filename
from docx import Document
from openpyxl import Workbook
import os
import tempfile
import uuid
import shutil

app = Flask(__name__)
app.secret_key = "secret-key"
ALLOWED_EXT = {"txt"}

# Temporary directory for converted files
TMP_DIR = tempfile.mkdtemp(prefix="converted_")

# ── HTML Interface ──
HTML_TEMPLATE = """
<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Text to DOCX & Excel</title>
  <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css"
        rel="stylesheet">
</head>
<body class="bg-light py-5">
<div class="container">
  <h1 class="text-center mb-4">Convert .txt File to DOCX and Excel</h1>

  {% with m = get_flashed_messages() %}
    {% if m %}
      <div class="alert alert-warning">{{ m[0] }}</div>
    {% endif %}
  {% endwith %}

  <form method="POST" enctype="multipart/form-data" class="card p-4 shadow-sm">
    <div class="mb-3">
      <input class="form-control" type="file" name="textfile" accept=".txt" required>
    </div>
    <button class="btn btn-primary">Convert</button>
  </form>

  {% if fid %}
  <div class="card mt-4 p-3 shadow-sm">
    <h5>Converted Files:</h5>
    <a class="btn btn-outline-success me-2" href="{{ url_for('download_file', fid=fid, ext='docx') }}">Download DOCX</a>
    <a class="btn btn-outline-success" href="{{ url_for('download_file', fid=fid, ext='xlsx') }}">Download Excel</a>
  </div>
  {% endif %}
</div>
</body>
</html>
"""

# ── Routes ──
@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        file = request.files.get("textfile")
        if not file or file.filename == "":
            flash("No file selected.")
            return redirect(request.url)

        if not file.filename.lower().endswith(".txt"):
            flash("Only .txt files are allowed.")
            return redirect(request.url)

        try:
            text = file.read().decode("utf-8")
        except Exception:
            flash("Could not read the text file.")
            return redirect(request.url)

        # Generate unique ID for this session
        fid = str(uuid.uuid4())
        docx_path = os.path.join(TMP_DIR, f"{fid}.docx")
        xlsx_path = os.path.join(TMP_DIR, f"{fid}.xlsx")

        # ---- Save DOCX ----
        doc = Document()
        doc.add_paragraph(text)
        doc.save(docx_path)

        # ---- Save Excel (line by line in column A) ----
        wb = Workbook()
        ws = wb.active
        for i, line in enumerate(text.splitlines(), start=1):
            ws.cell(row=i, column=1, value=line)
        wb.save(xlsx_path)

        return render_template_string(HTML_TEMPLATE, fid=fid)

    return render_template_string(HTML_TEMPLATE, fid=None)

@app.route("/download/<fid>.<ext>")
def download_file(fid, ext):
    if ext not in {"docx", "xlsx"}:
        flash("Invalid file type.")
        return redirect(url_for("index"))

    path = os.path.join(TMP_DIR, f"{fid}.{ext}")
    if not os.path.exists(path):
        flash("File not found or expired.")
        return redirect(url_for("index"))

    mimetype = {
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    }[ext]

    return send_file(path, mimetype=mimetype, as_attachment=True, download_name=f"converted.{ext}")

# Clean up temp files when app exits
import atexit
@atexit.register
def cleanup_temp_dir():
    shutil.rmtree(TMP_DIR, ignore_errors=True)

# ── Start the server ──
if __name__ == "__main__":
    app.run(debug=True)
